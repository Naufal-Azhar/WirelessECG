import os
import random
import numpy as np
from pathlib import Path
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_ecg_plot_image(
    output_path: str,
    recorded_processed: list,
    sample_rate: int = 100,
) -> bool:
    """Generate ECG plot image for the report.
    Port of ECGWindow.create_ecg_plot_image() - identical logic."""
    if not recorded_processed:
        return False

    proc_data_list = recorded_processed
    full_signal_data = np.array([row[1:] for row in proc_data_list], dtype=float)
    total_samples = len(full_signal_data)

    target_duration = 10
    required_samples = int(target_duration * sample_rate)

    if total_samples > required_samples:
        max_start_index = total_samples - required_samples
        start_index = random.randint(0, max_start_index)
        end_index = start_index + required_samples
        signal_data = full_signal_data[start_index:end_index]
        t = np.linspace(0, target_duration, required_samples)
        start_time_sec = start_index / sample_rate
        end_time_sec = end_index / sample_rate
        time_info = f"(Segmen {start_time_sec:.1f}s - {end_time_sec:.1f}s)"
    else:
        signal_data = full_signal_data
        duration_sec = total_samples / sample_rate
        t = np.linspace(0, duration_sec, total_samples)
        time_info = "(Full Record)"

    if len(t) != len(signal_data):
        min_len = min(len(t), len(signal_data))
        t = t[:min_len]
        signal_data = signal_data[:min_len]

    plt.close("all")
    fig, axes = plt.subplots(3, 1, figsize=(10, 6), sharex=True, dpi=100)
    fig.subplots_adjust(hspace=0.4)

    leads_label = ["Lead I", "Lead II", "Lead III"]
    colors = ["black", "black", "black"]

    for i, ax in enumerate(axes):
        ax.plot(t, signal_data[:, i], color=colors[i], linewidth=0.8)
        title_text = f"{leads_label[i]} {time_info if i == 0 else ''}"
        ax.set_title(title_text, loc="left", fontsize=10, fontweight="bold")

        ax.set_facecolor("white")
        ax.xaxis.set_major_locator(ticker.MultipleLocator(0.5))
        ax.xaxis.set_minor_locator(ticker.MultipleLocator(0.1))
        ax.yaxis.set_major_locator(ticker.MaxNLocator(nbins=6))

        ax.grid(which="major", axis="x", color="#ff9999", linestyle="-", linewidth=0.8)
        ax.grid(which="minor", axis="x", color="#ffcccc", linestyle="-", linewidth=0.4)
        ax.grid(which="major", axis="y", color="#ff9999", linestyle="-", linewidth=0.8)

        ax.tick_params(which="both", colors="black", labelbottom=False, labelleft=True)
        ax.ticklabel_format(useOffset=False, style="plain", axis="y")

        if len(signal_data[:, i]) > 0:
            y_min, y_max = np.min(signal_data[:, i]), np.max(signal_data[:, i])
            margin = 1.0 if y_min == y_max else (y_max - y_min) * 0.1
            ax.set_ylim(y_min - margin, y_max + margin)

    axes[-1].tick_params(labelbottom=True)
    axes[-1].set_xlabel("Time (seconds)")

    try:
        plt.savefig(output_path, bbox_inches="tight", pad_inches=0.1)
    except Exception as e:
        print(f"Plotting Error: {e}")
        plt.close(fig)
        return False

    plt.close(fig)
    return True


def generate_word_report(
    output_path: str,
    patient_name: str,
    patient_dob: str,
    heart_rate: int,
    cardiac_status: str,
    ai_results: list,
    recorded_processed: list,
    sample_rate: int = 100,
    duration_str: str = "00:00:00",
    total_raw_samples: int = 0,
):
    """Generate Word report. Port of ECGWindow.save_word_report() - identical logic."""
    document = Document()

    style = document.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    section = document.sections[0]
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)
    section.top_margin = Cm(1.27)
    section.bottom_margin = Cm(1.27)

    heading = document.add_heading("ECG MEDICAL REPORT", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph()

    # Patient Info Table
    table = document.add_table(rows=4, cols=4)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for row in table.rows:
        row.cells[0].width = Cm(3.0)
        row.cells[1].width = Cm(5.0)
        row.cells[2].width = Cm(3.2)
        row.cells[3].width = Cm(7.0)

    record_date_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    total_ai_segments = len(ai_results)
    jumlah_apnea = sum(1 for item in ai_results if item["status"] == "APNEA")
    kalimat_ai = f"{jumlah_apnea} dari {total_ai_segments} terdeteksi apnea"

    row0 = table.rows[0]
    row0.cells[0].text = "Patient Name"
    row0.cells[0].paragraphs[0].runs[0].bold = True
    row0.cells[1].text = f": {patient_name}"
    row0.cells[2].text = "Avg Heart Rate"
    row0.cells[2].paragraphs[0].runs[0].bold = True
    row0.cells[3].text = f": {heart_rate} BPM"

    row1 = table.rows[1]
    row1.cells[0].text = "Date of Birth"
    row1.cells[0].paragraphs[0].runs[0].bold = True
    row1.cells[1].text = f": {patient_dob}"
    row1.cells[2].text = "Cardiac Status"
    row1.cells[2].paragraphs[0].runs[0].bold = True
    row1.cells[3].text = f": {cardiac_status}"

    row2 = table.rows[2]
    row2.cells[0].text = "Record Date"
    row2.cells[0].paragraphs[0].runs[0].bold = True
    row2.cells[1].text = f": {record_date_str}"
    row2.cells[2].text = "AI Result"
    row2.cells[2].paragraphs[0].runs[0].bold = True
    row2.cells[3].text = f": {kalimat_ai}"

    row3 = table.rows[3]
    row3.cells[0].text = "Duration"
    row3.cells[0].paragraphs[0].runs[0].bold = True
    row3.cells[1].text = f": {duration_str}"
    row3.cells[2].text = "Total Raw Samples"
    row3.cells[2].paragraphs[0].runs[0].bold = True
    row3.cells[3].text = f": {total_raw_samples}"

    p_line = document.add_paragraph("_" * 95)
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph().paragraph_format.space_after = Pt(12)

    # Insert Plot Image
    temp_img_path = str(Path(output_path).parent / "temp_ecg_plot.png")

    if create_ecg_plot_image(temp_img_path, recorded_processed, sample_rate):
        h2 = document.add_heading("ECG Signal Visualization", level=2)
        h2.alignment = WD_ALIGN_PARAGRAPH.LEFT

        document.add_picture(temp_img_path, width=Cm(18.0))
        last_p = document.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        try:
            os.remove(temp_img_path)
        except Exception:
            pass
    else:
        document.add_paragraph("[No Signal Data Recorded]")

    document.save(str(output_path))
